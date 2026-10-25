"""Exportacao do relatorio: Excel formatado, PDF com tabela, links do WhatsApp com arquivo."""

import base64
import re
import uuid
import zlib
from datetime import date, time, timedelta
from io import BytesIO

from django.test import TestCase, override_settings
from django.urls import reverse
from openpyxl import load_workbook

from accounts.models import User

from .exports import build_pdf, build_xlsx, safe_filename
from .models import ProductionEntry, ProductionSheet

BASE = dict(ALLOWED_HOSTS=["testserver", "localhost"], SECURE_SSL_REDIRECT=False)
G, I = ProductionEntry.Destination.GUINCHO, ProductionEntry.Destination.IMPLEMENTO


def make_user(email):
    return User.objects.create_user(username=email, email=email, password="x-Senha-123456", role=User.Role.FUNCIONARIO)


def minutes(value):
    """Valor da coluna Horas em minutos (openpyxl devolve timedelta, time ou fracao do dia)."""
    if isinstance(value, timedelta):
        return round(value.total_seconds() / 60)
    if isinstance(value, time):
        return value.hour * 60 + value.minute
    return round(value * 1440)


def pdf_text(content):
    """Junta o texto de todas as paginas do PDF (streams do reportlab vem compactados)."""
    chunks = []
    for raw in re.findall(rb"stream\r?\n(.*?)endstream", content, re.S):
        data = raw.strip()
        if data.endswith(b"~>"):  # ASCII85 antes do Flate (padrao do reportlab)
            data = base64.a85decode(data, adobe=True)
        try:
            chunks.append(zlib.decompress(data).decode("latin-1"))
        except zlib.error:
            chunks.append(data.decode("latin-1", "ignore"))
    return "\n".join(chunks)


class ExportFixture(TestCase):
    def setUp(self):
        self.user = make_user("export@example.com")
        self.sheet = ProductionSheet.objects.create(
            owner=self.user, equipment="Chenlong 2", company_name="Empresa X", operator_name="Fagner",
            supervisor_name="Carlos",
        )
        e = lambda **kw: ProductionEntry.objects.create(sheet=self.sheet, **kw)
        e(work_date=date(2026, 10, 20), item_code="0011-024", quantity=12, destination=I, start_time=time(7, 0), end_time=time(11, 0))
        e(work_date=date(2026, 10, 21), item_code="0011-031", destination=G, start_time=time(13, 0), end_time=time(14, 30))
        e(work_date=date(2026, 10, 21), destination=I, manual_minutes=90, note="Ajuste & <teste> com acentuação")
        e(work_date=date(2026, 10, 22), item_code="X 123", start_time=time(22, 0), end_time=time(1, 0))  # virou a noite
        e(work_date=date(2026, 10, 22), item_code="0011-050", destination=G, start_time=time(15, 0))  # sem fim


class XlsxExportTests(ExportFixture):
    def _book(self):
        return load_workbook(BytesIO(build_xlsx(self.sheet)))

    def _header_row(self, ws):
        for row in ws.iter_rows(min_col=1, max_col=1):
            if row[0].value == "Data":
                return row[0].row
        self.fail("cabecalho nao encontrado")

    def test_each_value_is_in_its_own_column(self):
        ws = self._book().active
        h = self._header_row(ws)

        self.assertEqual([c.value for c in ws[h]][:8], ["Data", "Código", "Qtd.", "Destino", "Início", "Fim", "Horas", "Observação"])
        first = [c.value for c in ws[h + 1]][:8]
        self.assertEqual(first[0].date(), date(2026, 10, 20))
        self.assertEqual(first[1:4], ["0011-024", 12, "Implemento"])
        self.assertEqual((first[4], first[5]), (time(7, 0), time(11, 0)))
        self.assertEqual(minutes(first[6]), 240)

    def test_cells_use_real_excel_formats(self):
        ws = self._book().active
        h = self._header_row(ws)

        self.assertEqual(ws.cell(h + 1, 1).number_format, "DD/MM/YYYY")
        self.assertEqual(ws.cell(h + 1, 5).number_format, "HH:MM")
        self.assertEqual(ws.cell(h + 1, 7).number_format, "[H]:MM")
        self.assertEqual(ws.cell(h, 1).font.name, "Arial")
        self.assertTrue(ws.cell(h, 1).font.bold)
        self.assertEqual(ws.freeze_panes, f"A{h + 1}")
        self.assertEqual(ws.auto_filter.ref, f"A{h}:H{h + 5}")
        self.assertGreaterEqual(ws.column_dimensions["H"].width, 30)

    def test_totals_are_formulas_over_exactly_the_data_rows(self):
        ws = self._book().active
        h = self._header_row(ws)
        first, last = h + 1, h + 5
        formulas = {ws.cell(r, 6).value: ws.cell(r, 7).value for r in range(last + 1, last + 7) if ws.cell(r, 6).value}

        self.assertEqual(formulas["Total de horas"], f"=SUM(G{first}:G{last})")
        self.assertEqual(formulas["Guincho"], f'=SUMIF(D{first}:D{last},"Guincho",G{first}:G{last})')
        self.assertEqual(formulas["Implemento"], f'=SUMIF(D{first}:D{last},"Implemento",G{first}:G{last})')
        self.assertEqual(formulas["Peças produzidas"], f"=SUM(C{first}:C{last})")
        # O que as formulas vao somar bate com o total do app (o "sem fim" e texto, fica fora).
        hours = [ws.cell(r, 7).value for r in range(first, last + 1)]
        numeric = sum(minutes(v) for v in hours if v != "sem fim")
        self.assertEqual(numeric, self.sheet.totals()["total"])
        self.assertIn("sem fim", hours)

    def test_overnight_and_info_block(self):
        ws = self._book().active
        h = self._header_row(ws)
        info = {ws.cell(r, 1).value: ws.cell(r, 2).value for r in range(3, h)}

        overnight = next(r for r in range(h + 1, h + 6) if ws.cell(r, 2).value == "X 123")
        self.assertEqual(minutes(ws.cell(overnight, 7).value), 180)  # 22:00 -> 01:00
        self.assertEqual(info["Máquina"], "Chenlong 2")
        self.assertEqual(info["Encarregado"], "Carlos")
        self.assertEqual(info["Período"], "20/10/2026 a 22/10/2026")
        self.assertIn("Chenlong 2", ws["A1"].value)

    def test_empty_report_still_builds(self):
        empty = ProductionSheet.objects.create(owner=self.user, equipment="Vazia")

        ws = load_workbook(BytesIO(build_xlsx(empty))).active

        self.assertIn("Vazia", ws["A1"].value)


class PdfExportTests(ExportFixture):
    def test_pdf_has_table_rows_and_totals(self):
        content = build_pdf(self.sheet)
        text = pdf_text(content)

        self.assertTrue(content.startswith(b"%PDF"))
        for expected in ("0011-024", "20/10/2026", "07:00", "11:00", "4:00", "Guincho", "Implemento", "sem fim", "10:00", "Carlos",
                         "(12)", "produzidas", "QTD."):
            self.assertIn(expected, text, expected)

    def test_special_characters_do_not_break_the_pdf(self):
        text = pdf_text(build_pdf(self.sheet))

        # O reportlab quebra o texto em pedacos; juntando, fica "Ajuste & <teste> com acentuacao".
        joined = "".join(re.findall(r"\(((?:[^()\\]|\\.)*)\) Tj", text))
        self.assertIn("Ajuste & <teste> com acentua\\347\\343o", joined)

    def test_long_report_spans_pages_and_repeats_the_header(self):
        for i in range(70):
            ProductionEntry.objects.create(sheet=self.sheet, work_date=date(2026, 10, 1) + timedelta(days=i % 20),
                                           item_code=f"L-{i:03d}", destination=G, manual_minutes=30)

        content = build_pdf(self.sheet)

        pages = len(re.findall(rb"/Type /Page\b", content))
        self.assertGreaterEqual(pages, 2)
        # cabecalho repetido em toda pagina com tabela (a ultima pode ter so os totais)
        self.assertGreaterEqual(pdf_text(content).count("(OBSERVA"), pages - 1)
        self.assertIn("L-069", pdf_text(content))

    def test_empty_report_pdf(self):
        empty = ProductionSheet.objects.create(owner=self.user, equipment="Vazia")

        self.assertIn("Nenhum item anotado", pdf_text(build_pdf(empty)))

    def test_filename_is_safe(self):
        odd = ProductionSheet.objects.create(owner=self.user, equipment='Máq "1"/2')

        self.assertRegex(safe_filename(odd, "pdf"), r'^relatorio-[\w-]+\.pdf$')


@override_settings(**BASE)
class ExportViewsTests(ExportFixture):
    def test_owner_downloads_excel_and_pdf(self):
        self.client.force_login(self.user)

        xlsx = self.client.get(reverse("production_sheet_xlsx", args=[self.sheet.id]))
        pdf = self.client.get(reverse("production_sheet_pdf", args=[self.sheet.id]))
        view = self.client.get(reverse("production_sheet_pdf", args=[self.sheet.id]), {"ver": "1"})

        self.assertEqual(xlsx["Content-Type"], "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
        self.assertIn('attachment; filename="relatorio-Chenlong-2.xlsx"', xlsx["Content-Disposition"])
        self.assertEqual(pdf["Content-Type"], "application/pdf")
        self.assertTrue(pdf.content.startswith(b"%PDF"))
        self.assertTrue(view["Content-Disposition"].startswith("inline"))

    def test_other_user_and_anonymous_cannot_download_by_id(self):
        self.client.force_login(make_user("outro@example.com"))
        for name in ("production_sheet_xlsx", "production_sheet_pdf"):
            self.assertEqual(self.client.get(reverse(name, args=[self.sheet.id])).status_code, 404)
        self.client.logout()
        self.assertEqual(self.client.get(reverse("production_sheet_pdf", args=[self.sheet.id])).status_code, 302)

    def test_secret_link_opens_without_login_and_wrong_code_is_404(self):
        pdf = self.client.get(reverse("production_public_pdf", args=[self.sheet.share_token]))
        xlsx = self.client.get(reverse("production_public_xlsx", args=[self.sheet.share_token]))
        wrong = self.client.get(reverse("production_public_pdf", args=[uuid.uuid4()]))

        self.assertEqual(pdf.status_code, 200)
        self.assertTrue(pdf["Content-Disposition"].startswith("inline"))
        self.assertEqual(xlsx.status_code, 200)
        self.assertEqual(wrong.status_code, 404)

    def test_each_report_has_its_own_secret_code(self):
        other = ProductionSheet.objects.create(owner=self.user, equipment="Outra")

        self.assertNotEqual(other.share_token, self.sheet.share_token)

    def test_whatsapp_message_carries_pdf_and_excel_links(self):
        from urllib.parse import unquote

        self.client.force_login(self.user)
        location = self.client.get(reverse("production_sheet_whatsapp", args=[self.sheet.id]))["Location"]
        text = unquote(location.split("text=", 1)[1])

        self.assertIn(f"PDF: http://testserver/relatorio-producao/{self.sheet.share_token}/pdf/", text)
        self.assertIn(f"Excel: http://testserver/relatorio-producao/{self.sheet.share_token}/excel/", text)

    def test_detail_page_offers_pdf_share_and_downloads(self):
        self.client.force_login(self.user)

        html = self.client.get(reverse("production_sheet_detail", args=[self.sheet.id])).content.decode()

        self.assertIn("Enviar PDF pelo WhatsApp para Carlos", html)
        self.assertIn(f'data-share-pdf="{reverse("production_sheet_pdf", args=[self.sheet.id])}"', html)
        self.assertIn('data-share-name="relatorio-Chenlong-2.pdf"', html)
        self.assertIn("Baixar Excel", html)
        self.assertIn("js/production_share.js", html)
        self.assertNotIn("Exportar planilha (CSV)", html)
