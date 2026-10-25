"""Exportacao do relatorio de producao: planilha Excel (.xlsx) e PDF com tabela formatada.

As duas saidas usam as mesmas linhas (rows) para nunca divergirem entre si.
"""

from io import BytesIO

from django.utils import timezone

from .models import ProductionEntry, format_minutes

from accounts import print_theme as pt  # noqa: E402

# Mesmo visual dos PDFs (accounts/print_theme.py): tinta grafite, papel quente, sem azul.
INK = pt.INK
RULE = pt.RULE
PAPER = pt.PAPER
GUINCHO_FILL = pt.SAND
IMPLEMENTO_FILL = pt.SAGE
PENDING_FILL = pt.WARN_SOFT


def sheet_rows(sheet):
    """Linhas da tabela, ja na ordem de exibicao."""
    rows = []
    for e in sheet.entries.all():
        rows.append({
            "date": e.work_date,
            "code": e.item_code or "",
            "destination": e.get_destination_display() if e.destination else "",
            "destination_key": e.destination,
            "start": e.start_time,
            "end": e.end_time,
            "minutes": e.duration_minutes,
            "note": e.note or "",
        })
    return rows


def period_label(sheet):
    start, end = sheet.date_range()
    if not start:
        return "Sem itens"
    if start == end:
        return start.strftime("%d/%m/%Y")
    return f"{start:%d/%m/%Y} a {end:%d/%m/%Y}"


def info_pairs(sheet):
    pairs = [("Máquina", sheet.equipment)]
    if sheet.company_name:
        pairs.append(("Empresa", sheet.company_name))
    if sheet.operator_name:
        pairs.append(("Operador", sheet.operator_name))
    if sheet.supervisor_name:
        pairs.append(("Encarregado", sheet.supervisor_name))
    pairs.append(("Período", period_label(sheet)))
    pairs.append(("Situação", sheet.get_status_display()))
    return pairs


def safe_filename(sheet, ext):
    name = "".join(c if c.isalnum() or c in "-_" else "-" for c in sheet.equipment).strip("-") or "relatorio"
    return f"relatorio-{name}.{ext}"


# --------------------------------------------------------------------- Excel


def build_xlsx(sheet):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter
    from openpyxl.workbook.properties import CalcProperties

    wb = Workbook()
    ws = wb.active
    ws.title = "Relatório"
    base = Font(name="Arial", size=10)
    bold = Font(name="Arial", size=10, bold=True)
    thin = Side(style="thin", color=RULE)
    border = Border(bottom=thin)
    head_border = Border(bottom=Side(style="medium", color=INK))
    columns = [("Data", 12), ("Código", 16), ("Destino", 13), ("Início", 9), ("Fim", 9), ("Horas", 9), ("Observação", 42)]
    last_col = get_column_letter(len(columns))

    ws["A1"] = f"Relatório de produção - {sheet.equipment}"
    ws["A1"].font = Font(name="Georgia", size=16, bold=True, color=INK)
    ws.merge_cells(f"A1:{last_col}1")
    ws.row_dimensions[1].height = 24

    row = 3
    for label, value in info_pairs(sheet):
        ws.cell(row=row, column=1, value=label).font = Font(name="Arial", size=9, color=pt.MUTED)
        cell = ws.cell(row=row, column=2, value=value)
        cell.font = base
        ws.merge_cells(start_row=row, start_column=2, end_row=row, end_column=4)
        row += 1

    header_row = row + 1
    for idx, (title, width) in enumerate(columns, start=1):
        cell = ws.cell(row=header_row, column=idx, value=title)
        cell.font = Font(name="Arial", size=9, bold=True, color=pt.MUTED)
        cell.alignment = Alignment(horizontal="right" if title == "Horas" else "left", vertical="center")
        cell.border = head_border
        ws.column_dimensions[get_column_letter(idx)].width = width
    ws.row_dimensions[header_row].height = 20

    rows = sheet_rows(sheet)
    first_data = header_row + 1
    for i, item in enumerate(rows):
        r = first_data + i
        values = [
            item["date"], item["code"], item["destination"], item["start"], item["end"],
            item["minutes"] / 1440 if item["minutes"] is not None else "sem fim", item["note"],
        ]
        fill = None
        if item["minutes"] is None:
            fill = PatternFill("solid", fgColor=PENDING_FILL)
        elif i % 2:
            fill = PatternFill("solid", fgColor=PAPER)
        for c, value in enumerate(values, start=1):
            cell = ws.cell(row=r, column=c, value=value)
            cell.font = base
            cell.border = border
            cell.alignment = Alignment(vertical="center", horizontal="right" if c == 6 else "left",
                                       wrap_text=(c == 7))
            if fill:
                cell.fill = fill
        ws.cell(row=r, column=1).number_format = "DD/MM/YYYY"
        ws.cell(row=r, column=4).number_format = "HH:MM"
        ws.cell(row=r, column=5).number_format = "HH:MM"
        ws.cell(row=r, column=6).number_format = "[H]:MM"
        dest_key = item["destination_key"]
        if dest_key and item["minutes"] is not None:
            ws.cell(row=r, column=3).fill = PatternFill(
                "solid", fgColor=GUINCHO_FILL if dest_key == ProductionEntry.Destination.GUINCHO else IMPLEMENTO_FILL)
    last_data = first_data + max(len(rows), 1) - 1

    # Totais com formula: se alguem corrigir uma hora na planilha, o total acompanha.
    total_row = last_data + 2
    hours = f"F{first_data}:F{last_data}"
    dest = f"C{first_data}:C{last_data}"
    totals = [
        ("Total de horas", f"=SUM({hours})"),
        ("Guincho", f'=SUMIF({dest},"Guincho",{hours})'),
        ("Implemento", f'=SUMIF({dest},"Implemento",{hours})'),
    ]
    for offset, (label, formula) in enumerate(totals):
        r = total_row + offset
        label_cell = ws.cell(row=r, column=5, value=label)
        label_cell.font = bold
        label_cell.alignment = Alignment(horizontal="right")
        value_cell = ws.cell(row=r, column=6, value=formula)
        value_cell.number_format = "[H]:MM"
        value_cell.font = Font(name="Arial", size=11 if offset else 12, bold=True, color=pt.ACCENT if offset == 0 else INK)
        value_cell.alignment = Alignment(horizontal="right")
        value_cell.border = Border(top=Side(style="medium", color=INK)) if offset == 0 else border
        if offset == 0:
            value_cell.fill = PatternFill("solid", fgColor=pt.ACCENT_SOFT)
    pending = sum(1 for item in rows if item["minutes"] is None)
    note_row = total_row + len(totals) + 1
    if pending:
        ws.cell(row=note_row, column=1,
                value=f"{pending} item(ns) sem horário de fim ficaram fora do total.").font = Font(name="Arial", size=9, italic=True)
        note_row += 1
    generated = timezone.localtime().strftime("%d/%m/%Y %H:%M")
    ws.cell(row=note_row, column=1, value=f"Gerado pelo HoraCerta em {generated}.").font = Font(
        name="Arial", size=8, italic=True, color=pt.MUTED)

    ws.freeze_panes = ws.cell(row=first_data, column=1)
    if rows:
        ws.auto_filter.ref = f"A{header_row}:{last_col}{last_data}"
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_title_rows = f"{header_row}:{header_row}"
    wb.calculation = CalcProperties(fullCalcOnLoad=True)

    buffer = BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


# ----------------------------------------------------------------------- PDF


def build_pdf(sheet):
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.units import cm
    from reportlab.platypus import KeepTogether, Paragraph, SimpleDocTemplate, Spacer

    page = landscape(A4)
    buffer = BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=page, leftMargin=1.6 * cm, rightMargin=1.6 * cm, topMargin=1.5 * cm,
                            bottomMargin=1.7 * cm, title=f"Relatório de produção - {sheet.equipment}", author="HoraCerta")
    width = page[0] - doc.leftMargin - doc.rightMargin
    st = pt.styles()

    who = [p for p in (sheet.operator_name and f"Operador {pt.esc(sheet.operator_name)}",
                        sheet.company_name and pt.esc(sheet.company_name)) if p]
    story = pt.header_block(f"Relatório de produção · {sheet.equipment}", [
        " · ".join(who + [f"Período {pt.esc(period_label(sheet))}"]),
    ])
    story.append(pt.info_grid(info_pairs(sheet), width, columns=min(6, len(info_pairs(sheet)))))
    story.append(Spacer(1, 12))

    rows, fills = [], {}
    for i, item in enumerate(sheet_rows(sheet)):
        if item["minutes"] is None:
            fills[i] = PENDING_FILL
        rows.append([
            item["date"].strftime("%d/%m/%Y"),
            item["code"] or "-",
            item["destination"] or "-",
            item["start"].strftime("%H:%M") if item["start"] else "-",
            item["end"].strftime("%H:%M") if item["end"] else "-",
            item["note"],
            format_minutes(item["minutes"]) if item["minutes"] is not None else "sem fim",
        ])
    widths = [2.5, 3.4, 2.7, 1.8, 1.8]
    widths = [w * cm for w in widths] + [width - (sum(widths) + 2.2) * cm, 2.2 * cm]
    story.append(pt.data_table(["Data", "Código", "Destino", "Início", "Fim", "Observação", "Horas"], rows, widths,
                               numeric_cols=(6,), empty_text="Nenhum item anotado.", row_fills=fills))
    story.append(Spacer(1, 14))

    totals = sheet.totals()
    closing = [pt.totals_block(
        [("Guincho", format_minutes(totals[ProductionEntry.Destination.GUINCHO])),
         ("Implemento", format_minutes(totals[ProductionEntry.Destination.IMPLEMENTO]))],
        ("Total de horas", format_minutes(totals["total"])),
        width * 0.38,
    )]
    if totals["pending"]:
        closing += [Spacer(1, 4), Paragraph(
            f"{totals['pending']} item(ns) sem horário de fim ficaram fora do total.", st["small"])]
    story.append(KeepTogether(closing))

    decorate = pt.page_decorator(f"Relatório de produção · {sheet.equipment}")
    doc.build(story, onFirstPage=decorate, onLaterPages=decorate)
    return buffer.getvalue()
